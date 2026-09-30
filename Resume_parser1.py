"""
resume_parser1.py

Reads a resume file (PDF, DOCX, or TXT), pulls out the raw text, and uses
lightweight NLP / regex heuristics to extract the "essential information"
called for in the project brief: contact details, skills, and education.

This module purposefully has no dependency on BERT or Groq -- it is the
"input" stage of the pipeline used by resume_analyzer1.py.
"""

import os
import re

import pdfplumber
import docx


# ---------------------------------------------------------------------------
# Regex patterns for contact info
# ---------------------------------------------------------------------------
EMAIL_REGEX = r"[a-zA-Z0-9.\-_+]+@[a-zA-Z0-9\-]+\.[a-zA-Z0-9\-.]+"
PHONE_REGEX = r"(\+?\d{1,3}[\s.\-]?)?\(?\d{2,4}\)?[\s.\-]?\d{3,4}[\s.\-]?\d{3,4}"
LINKEDIN_REGEX = r"(https?://)?(www\.)?linkedin\.com/in/[A-Za-z0-9\-_/]+"
GITHUB_REGEX = r"(https?://)?(www\.)?github\.com/[A-Za-z0-9\-_/]+"


# ---------------------------------------------------------------------------
# Keyword lists used for simple keyword-spotting (NLP "extraction")
# ---------------------------------------------------------------------------
SKILL_KEYWORDS = [
    # Programming languages
    "Python", "Java", "C++", "C", "C#", "JavaScript", "TypeScript", "R", "Go",
    "Kotlin", "Swift", "PHP", "Scala", "MATLAB", "Solidity",
    # Web / frameworks
    "HTML", "CSS", "React", "Angular", "Vue", "Node.js", "Django", "Flask",
    "FastAPI", "Spring", ".NET", "Bootstrap",
    # Data / ML
    "Machine Learning", "Deep Learning", "NLP", "Computer Vision",
    "TensorFlow", "PyTorch", "Scikit-learn", "Keras", "Pandas", "NumPy",
    "OpenCV", "Hugging Face", "BERT", "LLM",
    # Data engineering / big data
    "SQL", "MySQL", "PostgreSQL", "MongoDB", "NoSQL", "Hadoop", "Spark",
    "Kafka", "ETL", "Airflow",
    # Cloud / DevOps
    "AWS", "Azure", "GCP", "Docker", "Kubernetes", "Jenkins", "CI/CD",
    "Linux", "Git", "Terraform",
    # Visualization / BI
    "Tableau", "Power BI", "Excel",
    # Testing
    "Selenium", "JUnit", "Automation Testing", "Manual Testing", "Postman",
    # Engineering tools
    "AutoCAD", "SolidWorks", "Revit", "ANSYS",
    # Business / management / soft skills
    "Project Management", "Agile", "Scrum", "SAP", "Salesforce",
    "Communication", "Leadership", "Stakeholder Management", "Negotiation",
    "Blockchain", "Cybersecurity", "Networking",
]

EDUCATION_KEYWORDS = [
    "b.tech", "be ", "b.e.", "btech", "bachelor", "b.sc", "bsc", "bca",
    "m.tech", "mtech", "m.e.", "master", "msc", "m.sc", "mca", "mba",
    "phd", "ph.d", "doctorate", "diploma", "high school", "intermediate",
    "12th", "10th", "university", "college", "institute",
]


class ResumeParser:
    """Extracts raw text and structured fields from a resume file."""

    def __init__(self, file_path):
        self.file_path = file_path
        self.raw_text = self._extract_text()

    # ------------------------------------------------------------------
    # Text extraction
    # ------------------------------------------------------------------
    def _extract_text(self):
        ext = os.path.splitext(self.file_path)[1].lower()
        if ext == ".pdf":
            return self._extract_from_pdf()
        if ext in (".docx", ".doc"):
            return self._extract_from_docx()
        if ext == ".txt":
            with open(self.file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        raise ValueError(f"Unsupported file type: '{ext}'. Use PDF, DOCX, or TXT.")

    def _extract_from_pdf(self):
        text_parts = []
        with pdfplumber.open(self.file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return "\n".join(text_parts)

    def _extract_from_docx(self):
        document = docx.Document(self.file_path)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        # Also pull text out of any tables (common in resume templates)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)

    # ------------------------------------------------------------------
    # Structured field extraction
    # ------------------------------------------------------------------
    def extract_contact_info(self):
        email = re.search(EMAIL_REGEX, self.raw_text)
        phone = re.search(PHONE_REGEX, self.raw_text)
        linkedin = re.search(LINKEDIN_REGEX, self.raw_text, re.IGNORECASE)
        github = re.search(GITHUB_REGEX, self.raw_text, re.IGNORECASE)
        return {
            "email": email.group(0) if email else None,
            "phone": phone.group(0).strip() if phone else None,
            "linkedin": linkedin.group(0) if linkedin else None,
            "github": github.group(0) if github else None,
        }

    def extract_skills(self, skill_list=None):
        """Keyword-spot skills from `skill_list` (defaults to SKILL_KEYWORDS)."""
        skill_list = skill_list or SKILL_KEYWORDS
        text_lower = self.raw_text.lower()
        found = []
        for skill in skill_list:
            pattern = r"(?<![a-zA-Z0-9])" + re.escape(skill.lower()) + r"(?![a-zA-Z0-9])"
            if re.search(pattern, text_lower):
                found.append(skill)
        return sorted(set(found))

    def extract_education(self):
        """Return resume lines that look like education entries."""
        education_lines = []
        for line in self.raw_text.splitlines():
            line_clean = line.strip()
            if not line_clean:
                continue
            line_lower = line_clean.lower()
            if any(keyword in line_lower for keyword in EDUCATION_KEYWORDS):
                education_lines.append(line_clean)
        return education_lines

    def get_summary(self):
        """Return a dict with raw text plus all extracted fields."""
        return {
            "raw_text": self.raw_text,
            "contact_info": self.extract_contact_info(),
            "skills": self.extract_skills(),
            "education": self.extract_education(),
            "word_count": len(self.raw_text.split()),
        }


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Usage: python resume_parser1.py <path-to-resume>")
    else:
        parser = ResumeParser(sys.argv[1])
        import json

        print(json.dumps(parser.get_summary(), indent=2))